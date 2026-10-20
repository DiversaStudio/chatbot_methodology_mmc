"""LLM label datasets: prompts, batching, synthetic dedupe, agreement, review workbook.

All synthetic: a fake client stands in for `claude -p`, texts are made up, embeddings are hand-built.
Nothing here needs the client's data, the network or the `claude` CLI.
"""
import json

import numpy as np
import pandas as pd
import pytest

from sami import categorize as cz
from sami import llm_labels as ll


def _work(n=12):
    """12 labelled messages per class (4 test, 8 dev) plus unlabelled ones, in 2-message conversations."""
    rows = []
    for c in ll.CLASSES:
        for i in range(n):
            rows.append({"text": f"{c[:5]} mensaje {i:02d}", "label": c, "split": "test" if i < 4 else "dev"})
    rows += [{"text": f"sin etiqueta {i}", "label": None, "split": "unlabeled"} for i in range(6)]
    w = pd.DataFrame(rows)
    w["row"] = np.arange(len(w))
    w["Conversation ID"] = [f"c{i // 2}" for i in range(len(w))]
    w["Message #"] = [i % 2 for i in range(len(w))]
    w["in_na_sheet"] = w["label"] == ll.NONE
    w["context"] = ""
    return w


class FakeClient:
    """Labels every message as `answer` (or per `fn`), counts calls, can fail or hit the limit first."""
    model = "fake"

    def __init__(self, answer=ll.NONE, fail=0, limits=0):
        self.answer, self.fail, self.limits, self.calls = answer, fail, limits, 0

    def __call__(self, system, user, schema):
        self.calls += 1
        if self.limits:
            self.limits -= 1
            raise ll.UsageLimit("limit")
        if self.fail:
            self.fail -= 1
            return {"result": {"items": []}, "usage": {}}
        if "items" in schema["properties"]:
            ids = [json.loads(l)["id"] for l in user.splitlines()[1:]]
            return {"result": {"items": [{"id": i, "c": self.answer} for i in ids]}, "usage": {}}
        n = int(user.split("Write ")[1].split()[0]) if "Write " in user else 25
        return {"result": {"messages": [f"mensaje {self.calls} {k}" for k in range(n)]}, "usage": {}}


def test_prepare_drops_scrubbed_and_builds_context():
    n = 60                                                   # enough rows for the stratified split
    df = pd.DataFrame({
        "row": range(n), "Conversation ID": [f"c{i // 4}" for i in range(n)], "Message #": [i % 4 for i in range(n)],
        "text": [f"mensaje {i}" for i in range(n)],
        "label": [cz.CATEGORIES[i % 2] for i in range(n)], "in_na_sheet": False})
    out = ll.prepare(df, scrub=lambda t: (None, False) if t == "mensaje 1" else (t, False))
    assert pd.isna(out.loc[1, "text"])                       # dropped: never sent
    assert out.loc[2, "context"] == "mensaje 0"              # the dropped message is not context
    assert out.loc[4, "context"] == ""                       # first message of its conversation
    assert out.loc[7, "context"] == "mensaje 4 | mensaje 5 | mensaje 6"      # nothing from another conversation
    assert set(out["split"]) <= {"test", "dev"} and (out["split"] == "test").any()


def test_label_prompt_uses_dev_only_and_all_classes():
    w = _work()
    p = ll.build_label_prompt(w)
    ll.assert_no_test_text(p, w)
    for c in cz.CATEGORIES:
        assert f"- {c}\n" in p
    assert "{" not in p.replace('{"', "")                    # every placeholder was filled


def test_leak_check_fires():
    w = _work()
    test_text = w.loc[w["split"] == "test", "text"].iloc[0]
    with pytest.raises(AssertionError):
        ll.assert_no_test_text(f"examples: {json.dumps(test_text)}", w)


def test_check_labels_rejects_bad_answers():
    ok = {"items": [{"id": 1, "c": ll.NONE}, {"id": 2, "c": "Empleabilidad"}]}
    ll.check_labels(ok, [1, 2])
    for bad, ids in [({"items": ok["items"][:1]}, [1, 2]),
                     ({"items": ok["items"] + [{"id": 3, "c": ll.NONE}]}, [1, 2]),
                     ({"items": [{"id": 1, "c": "inventada"}, {"id": 2, "c": ll.NONE}]}, [1, 2]),
                     ({"items": [{"id": 1, "c": ll.NONE}, {"id": 1, "c": ll.NONE}]}, [1, 2])]:
        with pytest.raises(ValueError):
            ll.check_labels(bad, ids)


def test_batches_resume_retry_and_wait_on_limit(tmp_path):
    w = _work().set_index("row")
    jobs = ll.label_jobs(w, "t_")
    assert len(jobs) == -(-len(w) // ll.LABEL_BATCH)
    waits = []
    c = FakeClient(fail=1, limits=1)
    s = ll.run_batches(jobs, tmp_path, c, "sys", ll.LABEL_SCHEMA, ll.check_labels, workers=1, sleep=waits.append)
    assert s["failed"] == 0 and s["ok"] == len(jobs) and waits == [ll.LIMIT_WAIT]
    calls = c.calls
    again = ll.run_batches(jobs, tmp_path, c, "sys", ll.LABEL_SCHEMA, ll.check_labels, workers=1)
    assert again["skip"] == len(jobs) and c.calls == calls          # finished batches are not asked again
    saved = json.loads((tmp_path / "t_0000.json").read_text())
    assert saved["model"] == "fake" and saved["prompt"] and saved["date"]
    assert len(ll.read_labels(tmp_path)) == len(w)


def test_persistent_bad_answer_leaves_failed_file(tmp_path):
    jobs = ll.label_jobs(_work().set_index("row").iloc[:5], "f_")
    s = ll.run_batches(jobs, tmp_path, FakeClient(fail=99), "sys", ll.LABEL_SCHEMA, ll.check_labels, workers=1)
    assert s["failed"] == 1 and (tmp_path / "f_0000.failed").exists()


def test_synth_jobs_cover_every_class_from_dev_seeds():
    w = _work()
    jobs = ll.synth_jobs(w, {c: 60 for c in ll.CLASSES}, "r0_")
    assert len(jobs) == len(ll.CLASSES) * 3                          # 25 + 25 + 10
    assert sum(j[2]["n"] for j in jobs) == 60 * len(ll.CLASSES)
    test_texts = set(w.loc[w["split"] == "test", "text"])
    assert not any(json.dumps(t) in j[1] for j in jobs for t in test_texts)      # seeds never come from the test set
    none_job = next(j for j in jobs if j[2]["cat"] == ll.NONE)
    assert "no need" in none_job[1]


def test_dedupe_exact_near_and_per_class_limit():
    syn = pd.DataFrame({"text": ["hola uno", "Hola  UNO", "real text", "otro a", "otro b", "cuarto"],
                        "category": ["A", "A", "A", "A", "A", "B"]})
    out = ll.dedupe_synth(syn, ["real text"])
    assert out["text"].tolist() == ["hola uno", "otro a", "otro b", "cuarto"]
    vec = {"hola uno": [1, 0], "otro a": [0.999, 0.04], "otro b": [0, 1], "cuarto": [1, 0]}   # cuarto: other class
    emb = lambda ts: np.array([np.array(vec[t]) / np.linalg.norm(vec[t]) for t in ts])
    near = ll.dedupe_synth(syn, ["real text"], embed=emb)
    assert near["text"].tolist() == ["hola uno", "otro b", "cuarto"]  # "otro a" ~ "hola uno" in class A only
    assert len(ll.dedupe_synth(syn, [], limit=1)) == 2                 # one per class


def test_build_synthetic_refills_to_350_per_class(tmp_path):
    w = _work()
    syn = ll.build_synthetic(w, tmp_path, FakeClient(), "sys", workers=1)
    assert set(syn["category"].value_counts()) == {ll.SYNTH_PER_CLASS}
    assert len(syn) == ll.SYNTH_PER_CLASS * len(ll.CLASSES)


def test_agreement_on_a_hand_made_confusion():
    r = ll.agreement(["Empleabilidad"] * 3 + [ll.NONE], ["Empleabilidad", "Empleabilidad", ll.NONE, ll.NONE])
    assert r["n"] == 4 and r["accuracy"] == 0.75 and r["share_ninguna"] == 0.5
    assert r["report"].loc["Empleabilidad", "recall"] == pytest.approx(2 / 3)


def test_review_workbook_layout(tmp_path):
    from openpyxl import load_workbook
    w = _work()
    pred = {int(r): ("Empleabilidad" if r % 3 == 0 else c) for r, c in zip(w["row"], w["label"].fillna(ll.NONE))}
    syn = pd.DataFrame({"text": ["a", "b"], "category": ["Emprendimiento", "Riesgos de protección"]})
    path = tmp_path / "out.xlsx"
    ll.write_review_workbook(path, w, pred, syn, {0: "Emprendimiento", 1: ll.NONE}, {"Mensajes": 1}, today="2026-10-16")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Guia_revision", "Etiquetas", "Sinteticos", "_listas"]
    assert wb["_listas"].sheet_state == "hidden"
    hdr = [c.value for c in wb["Etiquetas"][1]]
    assert hdr[:5] == ["Contexto (3 mensajes previos)", "Mensaje", "Etiqueta cliente", "Etiqueta LLM", "Validación humana"]
    assert len(wb["Etiquetas"].data_validations.dataValidation) == 2
    assert wb["Sinteticos"]["G2"].value == "sí"                       # the mismatch is listed first
    assert ll.MODEL in wb["Guia_revision"]["A3"].value


def test_module_imports_and_help_without_claude(capsys):
    with pytest.raises(SystemExit) as e:
        ll.main(["--help"])
    assert e.value.code == 0 and "prep" in capsys.readouterr().out


def test_labels_table_mirrors_first_pass_columns():
    w = _work()
    pred = {int(r): ll.NONE for r in w["row"]}
    t = ll.labels_table(w, pred)
    assert list(t.columns[:12]) == ["Contexto (3 mensajes previos)", "Mensaje", "Etiqueta cliente", "Etiqueta LLM",
                                    "Validación humana", "Categoría corregida", "Comentario", "Fuente", "split",
                                    "id", "Conversation ID", "Message #"]
    assert set(t["split"]) == {"prueba", "entrenamiento", "sin etiqueta"}
    assert (t.loc[t["Etiqueta cliente"] == ll.NONE, "Coincide cliente"] == "sí").all()
    assert (t.loc[t["Etiqueta cliente"].isna(), "Coincide cliente"] == "").all()


def test_labels_xlsx_has_filter_dropdowns_and_shading(tmp_path):
    from openpyxl import load_workbook
    w = _work()
    pred = {int(r): "Empleabilidad" for r in w["row"]}
    t = ll.labels_table(w, pred)
    ll.write_labels_xlsx(tmp_path / "t.xlsx", t)
    ws = load_workbook(tmp_path / "t.xlsx")["Etiquetas"]
    assert ws.auto_filter.ref.endswith(str(len(t) + 1)) and ws.freeze_panes == "A2"
    assert len(ws.data_validations.dataValidation) == 2
    assert ws.cell(2, 1).fill.fgColor.rgb == "00FFF2CC"


# ------------------------------------------------------------------ corrections (v2)
def _reviewed(tmp_path, w, fixes):
    """A reviewed workbook like the client's: `Etiquetas` with `id`, v1 Haiku label and `Categoría corregida`."""
    x = pd.DataFrame({"Mensaje": w["text"], "Etiqueta LLM": w["label"].fillna(ll.NONE), "id": w["row"],
                      "Categoría corregida": [fixes.get(int(r)) for r in w["row"]]})
    path = tmp_path / "reviewed.xlsx"
    x.to_excel(path, sheet_name="Etiquetas", index=False)
    return path


def _dev_rows(w, cls, k):
    return [int(r) for r in w.loc[(w["label"] == cls) & (w["split"] == "dev"), "row"][:k]]


def test_load_corrections_joins_and_classifies_verdicts(tmp_path):
    w = _work()
    a, b = _dev_rows(w, "Empleabilidad", 2)
    t = int(w.loc[(w["label"] == "Emprendimiento") & (w["split"] == "test"), "row"].iloc[0])
    path = _reviewed(tmp_path, w, {a: "Emprendimiento", b: "Empleabilidad", t: "Ninguna"})
    c = ll.load_corrections(path, w).set_index("row")
    assert len(c) == 3 and c.loc[t, "split"] == "test"
    assert c.loc[a, "verdict"] == "other" and c.loc[b, "verdict"] == "haiku" and c.loc[t, "verdict"] == "other"
    assert not c["conflict"].any()


def test_load_corrections_rejects_unknown_class_and_missing_id(tmp_path):
    w = _work()
    a = _dev_rows(w, "Empleabilidad", 1)[0]
    with pytest.raises(ValueError, match="unknown"):
        ll.load_corrections(_reviewed(tmp_path, w, {a: "Inventada"}), w)
    with pytest.raises(ValueError, match="not in the work"):
        ll.load_corrections(_reviewed(tmp_path, w, {a: "Ninguna"}), w.drop(index=a))


def test_conflicting_duplicate_text_is_flagged_and_never_a_shot(tmp_path):
    w = _work()
    a, b = _dev_rows(w, "Empleabilidad", 2)
    w.loc[w["row"] == b, "text"] = w.loc[w["row"] == a, "text"].iloc[0]
    c = ll.load_corrections(_reviewed(tmp_path, w, {a: "Emprendimiento", b: "Ninguna"}), w)
    assert c["conflict"].all()
    assert w.loc[w["row"] == a, "text"].iloc[0] not in ll.build_label_prompt(w, corrections=c)


def test_corrected_prompt_uses_corrected_labels_dev_only_and_never_leaks_test(tmp_path):
    w = _work()
    a = _dev_rows(w, "Empleabilidad", 1)[0]
    t = int(w.loc[(w["label"] == "Riesgos de protección") & (w["split"] == "test"), "row"].iloc[0])
    c = ll.load_corrections(_reviewed(tmp_path, w, {a: "Emprendimiento", t: "Ninguna"}), w)
    prompt = ll.build_label_prompt(w, corrections=c)
    text = json.dumps(w.loc[w["row"] == a, "text"].iloc[0], ensure_ascii=False)
    assert f"{text} -> Emprendimiento" in prompt and f"{text} -> Empleabilidad" not in prompt
    ll.assert_no_test_text(prompt, w)
    assert a in ll.shot_rows(w, c) and t not in ll.shot_rows(w, c)


def test_prompt_without_corrections_is_unchanged_by_the_refactor():
    w = _work()
    assert ll.build_label_prompt(w) == ll.build_label_prompt(w, corrections=None)
    ll.assert_no_test_text(ll.build_label_prompt(w), w)


def test_labels_table_final_column_prefers_the_correction(tmp_path):
    w = _work()
    a = _dev_rows(w, "Empleabilidad", 1)[0]
    c = ll.load_corrections(_reviewed(tmp_path, w, {a: "Emprendimiento"}), w)
    pred = {int(r): ll.NONE for r in w["row"]}
    plain, t = ll.labels_table(w, pred), ll.labels_table(w, pred, c)
    assert list(t.columns[:len(plain.columns)]) == list(plain.columns)
    assert (t["Etiqueta LLM"] == plain["Etiqueta LLM"]).all()
    r = t.set_index("id")
    assert r.loc[a, "Etiqueta final"] == "Emprendimiento" and r.loc[a, "Fuente final"] == "corrección humana"
    assert r.loc[a + 1, "Etiqueta final"] == ll.NONE
    assert r.loc[a + 1, "Fuente final"] == "LLM, discrepa (sin revisar)"          # client Empleabilidad, Haiku Ninguna
    none_row = int(w.loc[w["label"] == ll.NONE, "row"].iloc[0])
    assert r.loc[none_row, "Fuente final"] == "LLM, coincide con cliente"
    free = int(w.loc[w["label"].isna(), "row"].iloc[0])
    assert r.loc[free, "Fuente final"] == "LLM, sin etiqueta del cliente"
    assert set(t["Fuente final"]) == {"corrección humana", "LLM, discrepa (sin revisar)",
                                      "LLM, coincide con cliente", "LLM, sin etiqueta del cliente"}
    ll.write_labels_xlsx(tmp_path / "t.xlsx", t)


def test_gold_and_score_cover_corrected_and_agreed_rows(tmp_path):
    w = _work()
    a, b = _dev_rows(w, "Empleabilidad", 2)
    c = ll.load_corrections(_reviewed(tmp_path, w, {a: "Emprendimiento"}), w)
    v1 = {int(r): l for r, l in zip(w["row"], w["label"].fillna(ll.NONE))}
    v1[a] = "Emprendimiento"; v1[b] = "Ninguna"            # a: Haiku right vs client; b: unreviewed disagreement
    g = ll.gold_table(w, v1, c).set_index("row")
    assert g.loc[a, "gold"] == "Emprendimiento" and g.loc[a, "group"] == "other"
    assert b not in g.index                                  # no gold for unreviewed disagreements
    assert set(g["group"]) <= {"agreed", "haiku", "client", "other"}
    s = ll.score_gold(g.reset_index(), {int(r): g.loc[r, "gold"] for r in g.index}).set_index("subset")
    assert s.loc["all", "accuracy"] == 1.0


# ------------------------------------------------------------------ synthetic set sized from the final labels
def _final_work():
    """`_work()` with a final label: class sizes 12 (test 4) for all, with extras to make three classes the biggest."""
    w = _work()
    extra = pd.DataFrame([{"text": f"{c[:5]} extra {i}", "label": c, "split": "dev"}
                          for c, k in (("Empleabilidad", 9), ("Ninguna", 6), ("Emprendimiento", 3)) for i in range(k)])
    w = pd.concat([w, extra], ignore_index=True)
    w["row"] = np.arange(len(w))
    w["Conversation ID"] = [f"c{i // 2}" for i in range(len(w))]
    w["Message #"] = [i % 2 for i in range(len(w))]
    w["in_na_sheet"] = w["label"] == ll.NONE
    w["context"] = ""
    w["final"] = w["label"]
    w["fuente_final"] = np.where(w["row"] % 5 == 0, ll.FINAL_UNREVIEWED, "LLM, coincide con cliente")
    return w


def test_synth_targets_3x_for_the_three_biggest_10x_for_the_rest_without_test():
    wf = _final_work()
    t = ll.synth_targets(wf)
    real = wf[(wf["split"] != "test") & wf["final"].notna()]["final"].value_counts()
    big = set(real.index[:3])
    assert big == {"Empleabilidad", ll.NONE, "Emprendimiento"}
    assert all(t[c] == real[c] * (3 if c in big else 10) for c in ll.CLASSES)


def test_seed_pools_exclude_test_unreviewed_and_test_lookalikes():
    wf = _final_work()
    dup = wf[(wf["label"] == "Empleabilidad") & (wf["split"] == "dev")].index[0]
    test_text = wf.loc[(wf["label"] == "Empleabilidad") & (wf["split"] == "test"), "text"].iloc[0]
    wf.loc[dup, "text"] = test_text                                         # a dev/unlabelled repeat of a test text
    pools = ll.synth_pools(wf)
    test = set(wf.loc[wf["split"] == "test", "text"])
    assert not any(t in test for p in pools.values() for t in p)
    for c, p in pools.items():                                   # texts of this class's unreviewed rows are not seeds
        unreviewed = set(wf.loc[(wf["final"] == c) & (wf["fuente_final"] == ll.FINAL_UNREVIEWED), "text"]) - test
        assert not unreviewed & set(p)
    jobs = ll.synth_jobs(wf, {c: 25 for c in ll.CLASSES}, "r0_", pools=pools)      # the test guard does not fire
    assert len(jobs) == len(ll.CLASSES)


def test_synth_jobs_fail_when_a_prompt_holds_a_test_text():
    wf = _final_work()
    test_text = wf.loc[wf["split"] == "test", "text"].iloc[0]
    pools = {c: [test_text] * 3 for c in ll.CLASSES}
    with pytest.raises(AssertionError):
        ll.synth_jobs(wf, {c: 25 for c in ll.CLASSES}, "r0_", pools=pools)


def test_dedupe_accepts_per_class_limits():
    syn = pd.DataFrame({"text": ["a", "b", "c", "d", "e"], "category": ["A", "A", "A", "B", "B"]})
    assert ll.dedupe_synth(syn, [], limit={"A": 2, "B": 1})["category"].value_counts().to_dict() == {"A": 2, "B": 1}


def test_build_synthetic_reaches_per_class_targets_and_reports_a_shortfall(tmp_path):
    wf = _final_work()
    targets = {c: (30 if c == "Empleabilidad" else 0) for c in ll.CLASSES}
    targets["Riesgos de protección"] = 40
    pools = ll.synth_pools(wf)
    syn = ll.build_synthetic(wf, tmp_path, FakeClient(), "sys", targets=targets, pools=pools, workers=1)
    assert syn["category"].value_counts().to_dict() == {"Empleabilidad": 30, "Riesgos de protección": 40}

    class Repeats(FakeClient):                      # always the same messages: the class cannot reach its target
        def __call__(self, system, user, schema):
            n = int(user.split("Write ")[1].split()[0]) if "Write " in user else 25
            return {"result": {"messages": [f"igual {k}" for k in range(n)]}, "usage": {}}
    short = ll.build_synthetic(wf, tmp_path / "x", Repeats(), "sys", targets={"Emprendimiento": 50}, pools=pools, workers=1)
    assert len(short) == 25                         # not padded


def test_cached_embedder_embeds_each_text_once(tmp_path):
    calls = []
    def emb(texts):
        calls.append(list(texts))
        return np.array([[len(t), 1.0] for t in texts])
    f = ll.cached_embedder(tmp_path / "e.pkl", emb)
    f(["aa", "b"]); out = ll.cached_embedder(tmp_path / "e.pkl", emb)(["b", "aa", "ccc"])
    assert calls == [["aa", "b"], ["ccc"]] and out[:, 0].tolist() == [1, 2, 3]


def test_synthetic_workbook_layout(tmp_path):
    from openpyxl import load_workbook
    syn = pd.DataFrame({"text": ["uno", "dos", "tres"], "category": ["Emprendimiento", "Empleabilidad", "Ninguna"]})
    pred = {0: "Empleabilidad", 1: "Empleabilidad", 2: ll.NONE}                  # row 0 disagrees with its target
    table = pd.DataFrame({"Categoría": ["Emprendimiento"], "objetivo": [5], "generados": [3], "conservados": [1]})
    ll.write_synthetic_workbook(tmp_path / "s.xlsx", syn, pred, table, today="2026-10-19")
    wb = load_workbook(tmp_path / "s.xlsx")
    assert wb.sheetnames == ["Guia_revision", "Sinteticos", "_listas"] and wb["_listas"].sheet_state == "hidden"
    ws = wb["Sinteticos"]
    assert [c.value for c in ws[1]] == ["Mensaje sintético", "Categoría objetivo", "Etiqueta LLM", "Categoría corregida",
                                        "Comentario", "Revisar", "id"]
    assert ws["A2"].value == "uno" and ws["F2"].value == "sí" and ws["A3"].value == "dos"   # flagged first
    assert len(ws.data_validations.dataValidation) == 1 and ws.freeze_panes == "A2"
    assert any("faltan 4" in str(c.value) for c in wb["Guia_revision"]["A"])


def test_load_final_joins_on_id_and_rejects_bad_class(tmp_path):
    w = _work()
    x = pd.DataFrame({"id": w["row"], "Etiqueta final": w["label"].fillna(ll.NONE), "Fuente final": "LLM, coincide con cliente"})
    x.to_excel(tmp_path / "f.xlsx", index=False)
    out = ll.load_final(tmp_path / "f.xlsx", w)
    assert out["final"].notna().all() and "fuente_final" in out
    x.loc[0, "Etiqueta final"] = "Inventada"
    x.to_excel(tmp_path / "g.xlsx", index=False)
    with pytest.raises(ValueError, match="unknown"):
        ll.load_final(tmp_path / "g.xlsx", w)
