"""LLM-made datasets for the message classifier: labels for every message, plus synthetic messages.

Two datasets are built with Claude Haiku 5.5 from the client's original labels
(`datasets/labels/messages_category_original.xlsx`, read only):

* **Dataset 2, labelled messages.** Haiku labels every scrubbed message (development, the frozen
  test set labelled blind, and the ones the client never labelled). Agreement with the client on
  the test set says how far the model's labels can be trusted.
* **Dataset 3, synthetic messages.** 350 per class (the 10 categories and Ninguna), written from
  development examples only, deduplicated, then re-labelled by Haiku as a check.

Run it as a module (there is no other entry point)::

    PYTHONPATH=src .venv/bin/python -m sami.llm_labels {prep|pilot|label|synth|check|report|table} [--dry-run]
    ... -m sami.llm_labels --version v2 {corr|baseline|devcheck|testcheck|label|compare}
    ... -m sami.llm_labels --version synth {synth2|check2|synthreport} [--only CLASS] --workers 5
    ... -m sami.llm_labels final      # v1 labels + the reviewer's corrections -> messages_category_haiku_final.xlsx

Rules the numbers depend on, enforced in code:

* The frozen test set is `categorize.frozen_split`. Its client labels are never shown to the model
  and test texts never reach a prompt (`assert_no_test_text`).
* Only text scrubbed by `redact.scrub` is sent; messages it drops are never sent.
* Few-shot examples and synthetic seeds come from development rows only.
* Nothing prints or logs message text; the work folder (gitignored) holds scrubbed text only.
* Every batch answer is saved with the model id, prompt hash and date, so a rerun skips finished
  batches and the datasets can be rebuilt without calling the model again.

The model is reached through the Claude Code CLI in headless mode (`claude -p`), i.e. through the
user's subscription: no API key and no extra dependency. See `LLMClient`.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import re
import shutil
import subprocess
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from . import categorize as cz
from . import datasets

MODEL = "claude-haiku-5-5"
NONE = cz.NONE
CLASSES = list(cz.CLASSES)
LABEL_BATCH = 40          # messages per labelling call
SYNTH_BATCH = 25          # synthetic messages per generation call
SYNTH_PER_CLASS = 350
SYNTH_BIG, SYNTH_X_BIG, SYNTH_X_SMALL = 3, 3, 10     # the 3 biggest classes get 3x their real messages, the rest 10x
SYNTH_ROUNDS = 3          # extra rounds to refill classes that lose messages to the dedupe
NEAR_DUP = 0.97           # e5 cosine above which two synthetic messages of a class are the same
SHOTS = 8                 # development examples per class shown to the labeller
SEEDS = 10                # development examples per synthetic call
CONTEXT = 3               # previous messages of the conversation given as context
RETRIES = 3
WORKERS = 3
LIMIT_WAIT = 900          # seconds to wait when the subscription limit is hit, then resume
LIMIT_WAITS = 24          # ... at most this many times per batch (6 hours)


def work_dir(version: str = "v1") -> Path:
    """`datasets/labels/llm_haiku/` (v1) or `llm_haiku_v2/`: batches, parquet, prompts, CSVs. Gitignored."""
    return datasets.DATASETS_DIR / "labels" / ("llm_haiku" if version == "v1" else f"llm_haiku_{version}")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).lower().strip())


# ------------------------------------------------------------------ prep
def prepare(df: pd.DataFrame, scrub=None) -> pd.DataFrame:
    """Scrub the messages, build their conversation context and attach the frozen split.

    `df` is `categorize.load_labeled()`. `scrub(text) -> (clean | None, changed)`; messages it
    drops get `text = NaN` and are never sent. Columns of the result: `row`, `Conversation ID`,
    `Message #`, `text` (scrubbed), `context` (previous messages, scrubbed), `label` (client's,
    NaN when unlabelled), `in_na_sheet`, `split` in {test, dev, unlabeled}.
    """
    out = df[["row", "Conversation ID", "Message #", "label", "in_na_sheet"]].copy()
    is_test = cz.frozen_split(df)
    out["split"] = np.where(is_test, "test", np.where(df["label"].notna(), "dev", "unlabeled"))
    scrub = scrub or (lambda t: (t, False))
    out["text"] = [scrub(t)[0] for t in df["text"]]
    out = out.sort_values(["Conversation ID", "Message #"], kind="stable")
    ctx = {}
    for _, g in out.groupby("Conversation ID", sort=False):
        texts = g["text"].tolist()
        for i, r in enumerate(g["row"]):
            ctx[r] = " | ".join(t for t in texts[max(0, i - CONTEXT):i] if isinstance(t, str))
    out["context"] = out["row"].map(ctx)
    return out.sort_values("row").reset_index(drop=True)


def client_label(work: pd.DataFrame) -> pd.Series:
    """The client's class per message: `label`, or Ninguna for the `Mensajes_NA` repeats."""
    return work["label"].where(work["label"].notna(), np.where(work["in_na_sheet"], NONE, None))


# ------------------------------------------------------------------ prompt templates
# Placeholders are filled by `build_label_prompt` / `synth_jobs`. The few-shot examples and seeds are real
# (scrubbed) client text, so only these templates live in the repo, never a filled prompt.
LABEL_TEMPLATE = """You label messages that migrants and refugees (mostly Venezuelans in Colombia, Peru and Ecuador) sent to a humanitarian WhatsApp chatbot. Messages are informal, phone-typed Spanish with typos. Personal data is redacted: tokens such as [nombre] are placeholders, ignore them.

Give each message exactly ONE category, copied exactly from this closed list:
{cats}
- {none}  (the message expresses no need or question of its own)

Rules:
1. Label the person's NEED, not words that happen to appear in the message.
2. Always {none}, whatever the context says: greetings, thanks, "si" / "no" / "ok" / "listo", a bare place name or "estoy en <lugar>" answering where they are, a bare number or menu choice, a lone redacted token. A place counts only when the message also states a need: "estoy en cali y necesito comida" -> Asistencia Humanitaria.
3. "contexto" holds the previous messages of the same conversation. Use it ONLY to resolve a follow-up question or request that refers back to the topic being discussed ("y como lo saco", "cuanto cuesta", "donde queda"). Label only "mensaje", never the context.
4. If two categories fit, choose the dominant one.
5. Entidades del Estado vs Acceso a servicios del Estado: Entidades = the person asks WHICH institution or WHERE to go (an office, address, phone, who handles their case) without naming a public service. Servicios = the person asks how to GET a public service (health, EPS, SISBEN, education, social programs, social security), including requirements and eligibility. Naming a public service AND asking where to get it ("donde queda la oficina del sisben") is Acceso a servicios del Estado.
6. Asistencia Humanitaria vs Acceso a servicios del Estado: Asistencia Humanitaria = an urgent or direct need for help NOW (food, shelter, money, kits, "estoy enfermo, necesito un medico") or help from humanitarian organisations (ACNUR, OIM, fundaciones). Servicios = access to State systems and their requirements (EPS, affiliation, SISBEN, "me atienden sin papeles?", school, subsidies).
7. WHERE + a concrete need takes the category of the need: where to get a document -> Acceso a documentos y regularización; where to get or renew the salvoconducto -> Acceso a la protección internacional y asilo; where to report a crime or abuse -> Riesgos de protección; where to get a public service -> Acceso a servicios del Estado. Entidades del Estado only when they ask for an institution with no concrete need.
8. How to GET papers (PPT, pasaporte, cédula de extranjería, permits) -> Acceso a documentos y regularización. Whether they can use a service WITH the papers they have ("me atienden con el pasaporte?") -> Acceso a servicios del Estado.
9. Validating or apostilling diplomas -> Acceso a documentos y regularización. Studying, enrolling children in school, SENA courses -> Acceso a servicios del Estado.

Examples labelled by the client's team:
{shots}

Reply through the structured output: one item per input message, with the same "id" and the chosen category in "c".
"""

SYNTH_TEMPLATE = """You write realistic synthetic training data for a text classifier. The real data are messages that migrants and refugees (mostly Venezuelans in Colombia, Peru and Ecuador) send to a humanitarian WhatsApp chatbot.

Write {n} NEW messages that {goal}.
Here are real messages of that kind, for style and topic only (do NOT copy or lightly paraphrase them):
{seeds}

Requirements:
- Informal, phone-typed Spanish as in the examples: mostly lowercase, missing accents, some typos, short, occasional regional words (Venezuelan or Colombian).
- Vary the sub-topic, the situation (family, children, women, elderly, work, location), the phrasing and the length. No two messages alike.
- {avoid}
- No real names, phone numbers, ID numbers or addresses.

Reply through the structured output: the list of messages in "messages".
"""


# ------------------------------------------------------------------ prompts
def _dev_texts(work: pd.DataFrame, cls: str) -> list[str]:
    d = work[(work["split"] == "dev") & (work["label"] == cls) & work["text"].notna()]
    return list(dict.fromkeys(d["text"]))


def load_corrections(path: Path, work: pd.DataFrame) -> pd.DataFrame:
    """The reviewer's `Categoría corregida` rows of a reviewed Haiku workbook, joined to `work`.

    One row per corrected message: `row`, `text`, `split`, `client`, `llm_v1` (Haiku's first label),
    `correction`, `verdict` in {haiku, client, other} and `conflict` (the same text was corrected to
    two different classes, which context can justify, so such texts are never used as few-shots).
    """
    x = pd.read_excel(path, sheet_name="Etiquetas")
    x = x[x["Categoría corregida"].notna()]
    unknown = set(x["Categoría corregida"]) - set(CLASSES)
    if unknown:
        raise ValueError(f"unknown corrected classes: {sorted(unknown)}")
    w = work.assign(client=client_label(work)).set_index("row")
    missing = set(x["id"]) - set(w.index)
    if missing:
        raise ValueError(f"{len(missing)} corrected ids are not in the work table")
    out = pd.DataFrame({"row": x["id"].astype(int).to_numpy(), "llm_v1": x["Etiqueta LLM"].to_numpy(),
                        "correction": x["Categoría corregida"].to_numpy()})
    out["text"] = w.loc[out["row"], "text"].to_numpy()
    out["split"] = w.loc[out["row"], "split"].to_numpy()
    out["client"] = w.loc[out["row"], "client"].to_numpy()
    out["verdict"] = np.select([out["correction"] == out["llm_v1"], out["correction"] == out["client"]],
                               ["haiku", "client"], "other")
    out["conflict"] = out.groupby("text")["correction"].transform("nunique").gt(1)
    return out.sort_values("row").reset_index(drop=True)


def _shots(work: pd.DataFrame, corrections: pd.DataFrame | None, seed: int,
           hard: bool = True) -> list[tuple[int, str, str]]:
    """The few-shots as (row, text, class). With `corrections`, dev rows carry the reviewer's class
    instead of the client's, and (when `hard`) dev messages Haiku v1 got wrong come first (at most `SHOTS // 2`
    per class)."""
    rng = random.Random(seed)
    shots = []
    if corrections is None:
        for c in CLASSES:
            pool = _dev_texts(work, c)
            shots += [(-1, t, c) for t in rng.sample(pool, min(SHOTS, len(pool)))]
    else:
        fix = corrections[(corrections["split"] == "dev") & ~corrections["conflict"]].set_index("row")["correction"]
        d = work[(work["split"] == "dev") & work["text"].notna() & work["label"].notna()].copy()
        d["cls"] = d["row"].map(fix).fillna(d["label"])
        bad = set(corrections.loc[(corrections["verdict"] != "haiku") & (corrections["split"] == "dev"), "row"]) if hard else set()
        banned = set(corrections.loc[corrections["conflict"], "text"])
        for c in CLASSES:
            k = d[(d["cls"] == c) & ~d["text"].isin(banned)].drop_duplicates("text")
            hard = k[k["row"].isin(bad)]
            hard = hard.sample(min(SHOTS // 2, len(hard)), random_state=seed)
            rest = k[~k["row"].isin(hard["row"])]
            rest = rest.sample(min(SHOTS - len(hard), len(rest)), random_state=seed)
            shots += [(int(r), t, c) for r, t in zip(pd.concat([hard, rest])["row"], pd.concat([hard, rest])["text"])]
    rng.shuffle(shots)
    return shots


def shot_rows(work: pd.DataFrame, corrections: pd.DataFrame, seed: int = 21, hard: bool = True) -> set[int]:
    """Message ids shown as few-shots in the corrected prompt; they must not be scored on themselves."""
    return {r for r, _, _ in _shots(work, corrections, seed, hard)}


def build_label_prompt(work: pd.DataFrame, seed: int = 21, corrections: pd.DataFrame | None = None,
                       hard: bool = True) -> str:
    """The labelling system prompt: closed class list, separation rules and dev-only few-shots."""
    shots = [f"- {json.dumps(t, ensure_ascii=False)} -> {c}" for _, t, c in _shots(work, corrections, seed, hard)]
    tpl = LABEL_TEMPLATE
    return (tpl.replace("{cats}", "\n".join(f"- {c}" for c in cz.CATEGORIES))
            .replace("{none}", NONE).replace("{shots}", "\n".join(shots)))


def assert_no_test_text(prompt: str, work: pd.DataFrame) -> None:
    """Fail if any frozen-test text appears in a prompt (it would leak the answer)."""
    texts = set(work.loc[work["split"] == "test", "text"].dropna())
    leaked = [t for t in texts if json.dumps(t, ensure_ascii=False) in prompt]      # prompts quote every text
    if leaked:
        raise AssertionError(f"{len(leaked)} frozen-test texts appear in the prompt")


LABEL_SCHEMA = {
    "type": "object",
    "properties": {"items": {"type": "array", "items": {
        "type": "object",
        "properties": {"id": {"type": "integer"}, "c": {"type": "string", "enum": CLASSES}},
        "required": ["id", "c"]}}},
    "required": ["items"],
}
SYNTH_SCHEMA = {
    "type": "object",
    "properties": {"messages": {"type": "array", "items": {"type": "string"}}},
    "required": ["messages"],
}


def label_jobs(rows: pd.DataFrame, prefix: str, with_context: bool = True) -> list[tuple]:
    """Batches of `LABEL_BATCH` messages: (name, user prompt, meta), `rows` indexed by message id."""
    jobs = []
    for i in range(0, len(rows), LABEL_BATCH):
        b = rows.iloc[i:i + LABEL_BATCH]
        items = "\n".join(
            json.dumps({"id": int(k), "contexto": (c or "") if with_context else "", "mensaje": t},
                       ensure_ascii=False)
            for k, c, t in zip(b.index, b["context"] if with_context else [""] * len(b), b["text"]))
        jobs.append((f"{prefix}{i // LABEL_BATCH:04d}", "Messages:\n" + items, [int(k) for k in b.index]))
    return jobs


def check_labels(result: dict, ids: list[int]) -> None:
    """Raise unless `result` has exactly one valid label for each of `ids`. Never echoes text."""
    items = result.get("items")
    if not isinstance(items, list):
        raise TypeError("no items")
    got = [int(o["id"]) for o in items]
    if len(got) != len(set(got)) or set(got) != set(ids):
        raise ValueError("id mismatch")
    if any(o["c"] not in CLASSES for o in items):
        raise ValueError("unknown label")


# ------------------------------------------------------------------ the model
class UsageLimit(RuntimeError):
    """The subscription's usage limit was reached; waiting and retrying is the right response."""


class LLMClient:
    """Haiku through `claude -p` (headless Claude Code, subscription login).

    Each call: system prompt from a file (so client examples never appear in `ps`), the user
    prompt on stdin, no tools, no session, no slash commands, no settings, and an empty working
    directory so no project instructions leak into the call. The answer is forced through
    `--json-schema` and read from `structured_output`.
    """

    def __init__(self, model: str = MODEL, timeout: int = 900):
        self.model, self.timeout = model, timeout
        self.cwd = tempfile.mkdtemp(prefix="sami_llm_")

    def __call__(self, system: str, user: str, schema: dict) -> dict:
        with tempfile.NamedTemporaryFile("w", suffix=".txt", dir=self.cwd, delete=False, encoding="utf-8") as f:
            f.write(system)
        try:
            r = subprocess.run(
                ["claude", "-p", "--model", self.model, "--output-format", "json", "--tools", "",
                 "--no-session-persistence", "--disable-slash-commands", "--setting-sources", "",
                 "--system-prompt-file", f.name, "--json-schema", json.dumps(schema)],
                input=user, capture_output=True, text=True, timeout=self.timeout, cwd=self.cwd, check=False)
        finally:
            Path(f.name).unlink(missing_ok=True)
        try:
            d = json.loads(r.stdout)
        except json.JSONDecodeError:
            raise RuntimeError(f"claude exited {r.returncode} without JSON") from None
        blob = f"{d.get('result', '')} {d.get('api_error_status', '')}".lower()
        if d.get("is_error"):
            if "limit" in blob or "429" in blob or "overloaded" in blob:
                raise UsageLimit(str(d.get("api_error_status") or "limit"))
            raise RuntimeError(f"model error (status {d.get('api_error_status')})")
        if not isinstance(d.get("structured_output"), dict):
            raise TypeError("no structured output")
        u = d.get("usage", {})
        return {"result": d["structured_output"],
                "usage": {"in": u.get("input_tokens"), "cache_in": u.get("cache_creation_input_tokens"),
                          "cache_read": u.get("cache_read_input_tokens"), "out": u.get("output_tokens"),
                          "ms": d.get("duration_ms")}}


def run_batches(jobs, folder: Path, client, system: str, schema: dict, check, *, workers: int = WORKERS,
                sleep=time.sleep) -> dict:
    """Run `jobs` through `client`, one JSON per finished batch in `folder` (finished ones are skipped).

    A batch that fails validation is retried `RETRIES` times, then leaves `<name>.failed`. A usage
    limit makes the batch wait `LIMIT_WAIT` seconds and try again, without using up a retry.
    """
    folder.mkdir(parents=True, exist_ok=True)
    sys_hash = hashlib.sha256(system.encode("utf-8")).hexdigest()[:12]
    lock, done = threading.Lock(), {"n": 0}

    def one(job):
        name, user, meta = job
        out, failed = folder / f"{name}.json", folder / f"{name}.failed"
        if out.exists():
            return "skip"
        err, tries, waits = None, 0, 0
        while tries < RETRIES:
            try:
                ans = client(system, user, schema)
                check(ans["result"], meta)
                out.write_text(json.dumps(
                    {"meta": meta, "model": getattr(client, "model", MODEL), "prompt": sys_hash,
                     "date": datetime.now().astimezone().date().isoformat(), "usage": ans.get("usage"), "result": ans["result"]},
                    ensure_ascii=False), encoding="utf-8")
                failed.unlink(missing_ok=True)
                with lock:
                    done["n"] += 1
                    print(f"  {name} ok ({done['n']} done this run)", flush=True)
                return "ok"
            except UsageLimit:
                waits += 1
                if waits > LIMIT_WAITS:
                    err = "usage limit"
                    break
                with lock:
                    print(f"  {name}: usage limit, waiting {LIMIT_WAIT // 60} min", flush=True)
                sleep(LIMIT_WAIT)
            except Exception as e:  # noqa: BLE001 - any malformed answer is retried
                err, tries = f"{type(e).__name__}: {e}"[:200], tries + 1
        failed.write_text(str(err), encoding="utf-8")
        return "failed"

    with ThreadPoolExecutor(workers) as ex:
        res = list(ex.map(one, jobs))
    summary = {k: res.count(k) for k in ("ok", "skip", "failed")}
    print(summary, flush=True)
    return summary


def read_labels(folder: Path) -> dict[int, str]:
    """Message id -> label from every finished batch of `folder`."""
    pred = {}
    for f in sorted(folder.glob("*.json")):
        pred.update({int(o["id"]): o["c"] for o in json.loads(f.read_text(encoding="utf-8"))["result"]["items"]})
    return pred


# ------------------------------------------------------------------ synthetic messages
def _goal_avoid(cls: str) -> tuple[str, str]:
    if cls == NONE:
        goal = ('express no need or question of their own: greetings, thanks, "si", "ok", "listo", a bare '
                'place name or number, a lone redacted token like [nombre]')
        return goal, "None of them may state a need or ask anything about papers, work, health, help or services."
    others = "; ".join(c for c in cz.CATEGORIES if c != cls)
    return (f'clearly belong to the category "{cls}"',
            f'Each message must unambiguously belong to "{cls}" and not to any of: {others}.')


def synth_jobs(work: pd.DataFrame, need: dict[str, int], prefix: str, seed: int = 3,
               pools: dict[str, list[str]] | None = None) -> list[tuple]:
    """Generation calls for `need` messages per class, each seeded with examples of the class.

    Seeds come from `pools[class]` when given, else from the development rows. Every filled prompt is
    checked: a frozen-test text in it raises.
    """
    rng = random.Random(seed)
    tpl = SYNTH_TEMPLATE
    jobs = []
    for c in CLASSES:
        pool = pools[c] if pools is not None else _dev_texts(work, c)
        for k in range(-(-need.get(c, 0) // SYNTH_BATCH)):
            n = min(SYNTH_BATCH, need[c] - k * SYNTH_BATCH)
            goal, avoid = _goal_avoid(c)
            seeds = "\n".join(f"- {json.dumps(t, ensure_ascii=False)}" for t in rng.sample(pool, min(SEEDS, len(pool))))
            p = (tpl.replace("{n}", str(n)).replace("{goal}", goal).replace("{avoid}", avoid)
                 .replace("{seeds}", seeds))
            assert_no_test_text(p, work)
            jobs.append((f"{prefix}{CLASSES.index(c):02d}_{k:03d}", p, {"cat": c, "n": n}))
    return jobs


def check_synth(result: dict, meta: dict) -> None:
    msgs = result.get("messages")
    if not isinstance(msgs, list) or not all(isinstance(s, str) and s.strip() for s in msgs):
        raise ValueError("malformed")
    if len(msgs) < meta["n"] * 0.8:
        raise ValueError("too few messages")


def collect_synth(folder: Path) -> pd.DataFrame:
    """Every synthetic message of the finished batches, in batch order (before any dedupe)."""
    rows = []
    for f in sorted(folder.glob("*.json")):
        j = json.loads(f.read_text(encoding="utf-8"))
        rows += [{"text": s.strip(), "category": j["meta"]["cat"], "batch": f.stem} for s in j["result"]["messages"]]
    return pd.DataFrame(rows, columns=["text", "category", "batch"])


def dedupe_synth(syn: pd.DataFrame, real_texts, embed=None, limit: int | dict[str, int] = SYNTH_PER_CLASS,
                 threshold: float = NEAR_DUP) -> pd.DataFrame:
    """Drop exact copies of real or earlier synthetic text, then near-duplicates within a class.

    Keeps at most `limit` per class (one number, or a dict per class), first come first kept. `embed(list[str]) -> L2-normalised
    array`; without it only exact duplicates are removed.
    """
    real = {_norm(t) for t in real_texts}
    seen, keep = set(), []
    for i, t in enumerate(syn["text"]):
        k = _norm(t)
        if k not in real and k not in seen:
            seen.add(k)
            keep.append(i)
    syn = syn.iloc[keep].reset_index(drop=True)
    if embed is not None and len(syn):
        X = np.asarray(embed(syn["text"].tolist()))
        survive = np.zeros(len(syn), bool)
        for cat, idx in syn.groupby("category").indices.items():
            cap = limit[cat] if isinstance(limit, dict) else limit
            kept = []
            for i in idx:
                if len(kept) >= cap:
                    break
                if not kept or float((X[kept] @ X[i]).max()) < threshold:
                    kept.append(i)
            survive[kept] = True
        syn = syn[survive].reset_index(drop=True)
    elif len(syn):
        cap = syn["category"].map(limit) if isinstance(limit, dict) else limit
        syn = syn[syn.groupby("category", sort=False).cumcount() < cap].reset_index(drop=True)
    return syn


def build_synthetic(work: pd.DataFrame, folder: Path, client, system_synth: str, embed=None, *,
                    targets: dict[str, int] | None = None, pools: dict[str, list[str]] | None = None,
                    **kw) -> pd.DataFrame:
    """Generate until every class has its target (`SYNTH_PER_CLASS` each by default) or `SYNTH_ROUNDS` refills passed.

    A class that cannot reach its target is reported, never padded.
    """
    targets = targets or {c: SYNTH_PER_CLASS for c in CLASSES}
    need = dict(targets)
    real = work["text"].dropna().tolist()
    for rnd in range(SYNTH_ROUNDS + 1):
        run_batches(synth_jobs(work, need, f"r{rnd}_", seed=3 + rnd, pools=pools), folder, client, system_synth,
                    SYNTH_SCHEMA, check_synth, **kw)
        syn = dedupe_synth(collect_synth(folder), real, embed, limit=targets)
        have = syn["category"].value_counts().to_dict()
        need = {c: max(0, targets.get(c, 0) - have.get(c, 0)) for c in CLASSES}
        if not any(need.values()):
            break
        print(f"round {rnd}: still short {({c: n for c, n in need.items() if n})}", flush=True)
    return syn


# ------------------------------------------------------------------ synthetic set sized from the final labels
FINAL_UNREVIEWED = "LLM, discrepa (sin revisar)"


def load_final(path: Path, work: pd.DataFrame) -> pd.DataFrame:
    """`work` plus `final` and `fuente_final` from the final labelled workbook, joined on `id` = `row`."""
    x = pd.read_excel(path)
    bad = set(x["Etiqueta final"]) - set(CLASSES)
    if bad:
        raise ValueError(f"unknown final classes: {sorted(bad)}")
    if not set(x["id"]) <= set(work["row"]):
        raise ValueError("final labels hold ids that are not in the work table")
    m = x.drop_duplicates("id").set_index("id")
    return work.assign(final=work["row"].map(m["Etiqueta final"]), fuente_final=work["row"].map(m["Fuente final"]))


def synth_targets(wf: pd.DataFrame) -> dict[str, int]:
    """Synthetic messages per class: 3x the real non-test messages for the 3 biggest classes, 10x for the rest."""
    real = wf[(wf["split"] != "test") & wf["final"].notna() & wf["text"].notna()]["final"].value_counts()
    big = set(real.index[:SYNTH_BIG])
    return {c: int(real.get(c, 0)) * (SYNTH_X_BIG if c in big else SYNTH_X_SMALL) for c in CLASSES}


def synth_pools(wf: pd.DataFrame) -> dict[str, list[str]]:
    """Seed texts per class: non-test messages with a final label, without the unreviewed client/Haiku
    disagreements and without any text a frozen-test message also has (an unlabelled repeat of a test text)."""
    test = {_norm(t) for t in wf.loc[wf["split"] == "test", "text"].dropna()}
    d = wf[(wf["split"] != "test") & wf["final"].notna() & wf["text"].notna() & (wf["fuente_final"] != FINAL_UNREVIEWED)]
    d = d[~d["text"].map(_norm).isin(test)]
    return {c: list(dict.fromkeys(d.loc[d["final"] == c, "text"])) for c in CLASSES}


def cached_embedder(path: Path, embed):
    """`embed` that remembers every text it has seen in `path`, so a refill round embeds only the new ones."""
    import pickle
    memo = pickle.loads(path.read_bytes()) if path.exists() else {}

    def run(texts):
        new = [t for t in dict.fromkeys(texts) if t not in memo]
        if new:
            memo.update(zip(new, np.asarray(embed(new))))
            path.write_bytes(pickle.dumps(memo))
        return np.array([memo[t] for t in texts])
    return run


SYNTH_SYSTEM = "You write realistic synthetic Spanish chat messages for a text classifier."


# ------------------------------------------------------------------ agreement
def agreement(y_client, y_llm) -> dict:
    """Accuracy, Cohen's kappa, share labelled Ninguna and the per-class report of `y_llm` vs `y_client`."""
    from sklearn.metrics import classification_report, cohen_kappa_score
    y_client, y_llm = pd.Series(y_client).reset_index(drop=True), pd.Series(y_llm).reset_index(drop=True)
    rep = pd.DataFrame(classification_report(y_client, y_llm, labels=CLASSES, output_dict=True,
                                             zero_division=0)).T.loc[CLASSES]
    return {"n": len(y_client), "accuracy": float((y_client == y_llm).mean()),
            "kappa": float(cohen_kappa_score(y_client, y_llm)),
            "share_ninguna": float((y_llm == NONE).mean()), "report": rep}


# ------------------------------------------------------------------ labels table
def labels_table(work: pd.DataFrame, pred: dict[int, str], corrections: pd.DataFrame | None = None) -> pd.DataFrame:
    """Every message Haiku labelled, one row each, in the order of the client's sheet.

    Columns: context, message, the client's label, Haiku's label, three blank reviewer columns,
    `Fuente`, `split`, the keys (`id`, `Conversation ID`, `Message #`) and `Coincide cliente`.
    With `corrections`, two more columns: `Etiqueta final` (the reviewer's class where there is one,
    else Haiku's) and `Fuente final` (corrección humana / LLM, coincide con cliente / LLM, discrepa
    (sin revisar) / LLM, sin etiqueta del cliente). Nothing else is overwritten.
    """
    w = work.assign(llm=work["row"].map(pred), client=client_label(work))
    w = w[w["text"].notna() & w["llm"].notna()]
    out = pd.DataFrame({
        "Contexto (3 mensajes previos)": w["context"].to_numpy(), "Mensaje": w["text"].to_numpy(),
        "Etiqueta cliente": w["client"].to_numpy(), "Etiqueta LLM": w["llm"].to_numpy(),
        "Validación humana": None, "Categoría corregida": None, "Comentario": None,
        "Fuente": np.select([w["client"].notna() & w["label"].isna(), w["client"].notna()],
                            ["cliente (Mensajes_NA)", "cliente"], "LLM"),
        "split": np.select([w["split"] == "test", w["client"].notna()], ["prueba", "entrenamiento"], "sin etiqueta"),
        "id": w["row"].to_numpy(), "Conversation ID": w["Conversation ID"].to_numpy(),
        "Message #": w["Message #"].to_numpy()})
    out["Coincide cliente"] = np.select(
        [w["client"].isna().to_numpy(), (w["client"] == w["llm"]).to_numpy()], ["", "sí"], "no")
    if corrections is not None:
        fix = out["id"].map(corrections.drop_duplicates("row").set_index("row")["correction"])
        out["Etiqueta final"] = fix.fillna(out["Etiqueta LLM"])
        out["Fuente final"] = np.select(
            [fix.notna(), out["Etiqueta cliente"].isna(), out["Coincide cliente"] == "sí"],
            ["corrección humana", "LLM, sin etiqueta del cliente", "LLM, coincide con cliente"],
            "LLM, discrepa (sin revisar)")
    return out


def gold_table(work: pd.DataFrame, pred: dict[int, str], corrections: pd.DataFrame) -> pd.DataFrame:
    """The measuring stick: `row`, `split`, `gold`, `group`.

    Gold = the reviewer's class for corrected messages (`group` haiku / client / other = who the
    reviewer agreed with), and the shared label for unreviewed messages where the client and `pred`
    agree (`group` agreed). Unreviewed disagreements and Haiku-only messages have no gold and are left out.
    """
    fix = corrections.drop_duplicates("row").set_index("row")
    w = work.assign(client=client_label(work), llm=work["row"].map(pred))
    w = w[w["text"].notna() & w["llm"].notna() & w["split"].isin(["dev", "test"])]
    agreed = w[(w["client"] == w["llm"]) & ~w["row"].isin(fix.index)]
    rev = w[w["row"].isin(fix.index)]
    return pd.concat([
        pd.DataFrame({"row": agreed["row"], "split": agreed["split"], "gold": agreed["client"], "group": "agreed"}),
        pd.DataFrame({"row": rev["row"], "split": rev["split"], "gold": rev["row"].map(fix["correction"]),
                      "group": rev["row"].map(fix["verdict"])})], ignore_index=True)


def score_gold(gold: pd.DataFrame, pred: dict[int, str]) -> pd.DataFrame:
    """Accuracy of `pred` against `gold`, overall and per `group`/`split`, over the messages it labelled."""
    g = gold[gold["row"].isin(pred)].assign(ok=lambda d: d["gold"] == d["row"].map(pred))
    rows = [{"subset": "all", "n": len(g), "accuracy": g["ok"].mean() if len(g) else np.nan}]
    for key in ("group", "split"):
        for k, d in g.groupby(key):
            rows.append({"subset": f"{key}={k}", "n": len(d), "accuracy": d["ok"].mean()})
    return pd.DataFrame(rows)


def write_labels_xlsx(path: Path, table: pd.DataFrame) -> None:
    """`labels_table` as a workbook: filters, frozen header, class dropdowns, rows that differ shaded.

    Yellow = Haiku differs from the client. The reviewer columns carry the same dropdowns as the review workbook (hidden `_listas` sheet).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Etiquetas"
    lists = wb.create_sheet("_listas")
    for col, vals in enumerate((CLASSES, ["cliente", "LLM", "otra (ver Categoría corregida)"]), 1):
        for i, v in enumerate(vals, 1):
            lists.cell(i, col, v)
    lists.sheet_state = "hidden"
    cols = list(table.columns)
    width = {"Contexto (3 mensajes previos)": 45, "Mensaje": 50, "Etiqueta cliente": 30, "Etiqueta LLM": 30,
             "Categoría corregida": 30, "Comentario": 30}
    for j, h in enumerate(cols, 1):
        ws.cell(1, j, h).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(j)].width = width.get(h, 16)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    no_client = cols.index("Coincide cliente")
    for i, row in enumerate(table.itertuples(index=False), 2):
        for j, x in enumerate(row, 1):
            if x is not None and not pd.isna(x) and x != "":
                ws.cell(i, j, x.item() if hasattr(x, "item") else x)
        if row[no_client] == "no":
            for j in range(1, len(cols) + 1):
                ws.cell(i, j).fill = yellow
    last = len(table) + 1
    for h, ref in (("Validación humana", "=_listas!$B$1:$B$3"), ("Categoría corregida", f"=_listas!$A$1:$A${len(CLASSES)}")):
        dv = DataValidation(type="list", formula1=ref, allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(cols.index(h) + 1)
        dv.add(f"{letter}2:{letter}{last}")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ review workbook
def write_review_workbook(path: Path, work: pd.DataFrame, pred: dict[int, str], syn: pd.DataFrame,
                          syn_pred: dict[int, str], summary: dict, today: str | None = None) -> None:
    """Spanish review workbook for the client's reviewers: `Guia_revision`, `Etiquetas`, `Sinteticos`.

    `Etiquetas` lists only what needs validating: first messages where the client and Haiku
    differ, then the ones only Haiku labelled, each block grouped by Haiku's label. `Sinteticos`
    lists the synthetic messages grouped by Haiku's label (whatever class they were meant for),
    flagged when Haiku's check disagrees with the target class.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    today = today or datetime.now().astimezone().date().isoformat()
    w = work.assign(llm=work["row"].map(pred), client=client_label(work))
    w = w[w["text"].notna() & w["llm"].notna()]
    differ = w["client"].notna() & (w["client"] != w["llm"])
    only_llm = w["client"].isna()
    block = np.where(differ, 0, np.where(only_llm, 1, 2))
    et = w[block < 2].assign(_b=block[block < 2]).sort_values(["_b", "llm", "row"], kind="stable")
    split = np.select([et["split"] == "test", et["client"].notna()], ["prueba", "entrenamiento"], "sin etiqueta")
    etiquetas = pd.DataFrame({
        "Contexto (3 mensajes previos)": et["context"].to_numpy(), "Mensaje": et["text"].to_numpy(),
        "Etiqueta cliente": et["client"].to_numpy(), "Etiqueta LLM": et["llm"].to_numpy(),
        "Validación humana": None, "Categoría corregida": None, "Comentario": None,
        "Fuente": np.where(et["client"].notna(), "cliente", "LLM"), "split": split,
        "id": et["row"].to_numpy(), "Conversation ID": et["Conversation ID"].to_numpy(),
        "Message #": et["Message #"].to_numpy()})
    s = syn.assign(llm=[syn_pred.get(i) for i in syn.index])
    s["Revisar"] = np.where(s["llm"].notna() & (s["llm"] != s["category"]), "sí", "")
    s = s.sort_values(["llm", "Revisar", "category"], ascending=[True, False, True], kind="stable")
    sinteticos = pd.DataFrame({
        "Mensaje sintético": s["text"].to_numpy(), "Categoría objetivo": s["category"].to_numpy(),
        "Etiqueta LLM": s["llm"].to_numpy(), "Veredicto": None, "Categoría corregida": None,
        "Comentario": None, "Revisar": s["Revisar"].to_numpy(), "id": s.index.to_numpy()})

    guia = [
        "REVISIÓN DE ETIQUETAS (LLM)",
        "Etiqueta cliente = etiqueta original del cliente (messages_category_original.xlsx, sin cambios).",
        f"Etiqueta LLM = {MODEL} ({today}) siguiendo las reglas de separación entre categorías.",
        "REVISION",
        "Solo los mensajes a validar, en dos bloques: primero los que cliente y LLM no coinciden, después los",
        "etiquetados solo por el LLM. En cada bloque, agrupados por Etiqueta LLM.",
        "Contexto = los 3 mensajes previos de la misma conversación (lo que vio el LLM). Los textos están anonimizados.",
        "Validación humana: cliente / LLM / otra (en los de solo LLM: LLM u otra). Si es otra, elegir la categoría en Categoría corregida.",
        "split = prueba: mensajes reservados para evaluar el modelo (no se usan para entrenar).",
        "Ninguna = el mensaje no expresa una necesidad propia (saludos, gracias, si/ok, solo un lugar o un número).",
        "SINTETICOS",
        "Mensajes generados por el LLM para balancear las categorías (350 por categoría); no son mensajes reales.",
        "Agrupados por Etiqueta LLM, sin importar la categoría para la que se generaron.",
        "Revisar = sí: al etiquetarlos, el LLM no los puso en la categoría para la que se generaron.",
        "Veredicto: mantener / mover a Categoría corregida / eliminar.",
        "RESUMEN", *[f"{k}: {v}" for k, v in summary.items()]]
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Guia_revision")
    for i, line in enumerate(guia, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 120
    lists = wb.create_sheet("_listas")
    for col, vals in enumerate((CLASSES, ["cliente", "LLM", "otra (ver Categoría corregida)"],
                                ["mantener", "mover a Categoría corregida", "eliminar"]), 1):
        for i, v in enumerate(vals, 1):
            lists.cell(i, col, v)
    lists.sheet_state = "hidden"
    ref = {"cat": f"=_listas!$A$1:$A${len(CLASSES)}", "val": "=_listas!$B$1:$B$3", "ver": "=_listas!$C$1:$C$3"}
    yellow = PatternFill("solid", fgColor="FFF2CC")

    def sheet(name, df, widths, drops, flag=None):
        ws = wb.create_sheet(name)
        for j, h in enumerate(df.columns, 1):
            ws.cell(1, j, h).font = Font(bold=True)
            ws.column_dimensions[ws.cell(1, j).column_letter].width = widths.get(h, 16)
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, x in enumerate(row, 1):
                if x is not None and not pd.isna(x) and x != "":
                    ws.cell(i, j, x.item() if hasattr(x, "item") else x)
            if flag is not None and row[list(df.columns).index(flag)] == "sí":
                for j in range(1, len(df.columns) + 1):
                    ws.cell(i, j).fill = yellow
        for h, key in drops.items():
            col = ws.cell(1, list(df.columns).index(h) + 1).column_letter
            dv = DataValidation(type="list", formula1=ref[key], allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(df) + 1}")
        ws.freeze_panes = "A2"

    sheet("Etiquetas", etiquetas, {"Contexto (3 mensajes previos)": 45, "Mensaje": 50, "Etiqueta cliente": 30,
                                   "Etiqueta LLM": 30, "Validación humana": 18, "Categoría corregida": 30,
                                   "Comentario": 30},
          {"Validación humana": "val", "Categoría corregida": "cat"})
    sheet("Sinteticos", sinteticos, {"Mensaje sintético": 60, "Categoría objetivo": 30, "Etiqueta LLM": 30,
                                     "Veredicto": 22, "Categoría corregida": 30, "Comentario": 30},
          {"Veredicto": "ver", "Categoría corregida": "cat"}, flag="Revisar")
    wb.move_sheet("_listas", offset=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ CLI
def write_synthetic_workbook(path: Path, syn: pd.DataFrame, syn_pred: dict[int, str], table: pd.DataFrame,
                             today: str | None = None) -> None:
    """The review workbook of the synthetic set: `Guia_revision`, `Sinteticos`, hidden `_listas`.

    `syn` has `text` and `category` (index = id), `syn_pred` the check label per id, `table` the
    target / generated / kept count per class. Flagged rows (check label differs from the target) come
    last and are shaded; `Categoría corregida` is a dropdown, and empty means keep.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    today = today or datetime.now().astimezone().date().isoformat()
    s = syn.assign(llm=syn.index.map(syn_pred))
    s["rev"] = np.where(s["llm"].notna() & (s["llm"] != s["category"]), "sí", "")
    s["_c"] = s["category"].map({c: i for i, c in enumerate(CLASSES)})
    s = s.assign(_f=(s["rev"] == "sí")).sort_values(["_f", "_c"], kind="stable")
    df = pd.DataFrame({"Mensaje sintético": s["text"].to_numpy(), "Categoría objetivo": s["category"].to_numpy(),
                       "Etiqueta LLM": s["llm"].to_numpy(), "Categoría corregida": None, "Comentario": None,
                       "Revisar": s["rev"].to_numpy(), "id": s.index.to_numpy()})
    guia = [
        "MENSAJES SINTÉTICOS (LLM)",
        f"Generados con {MODEL} ({today}) a partir de mensajes reales etiquetados (sin el conjunto de prueba). No son mensajes reales.",
        "Cantidad por categoría = 3 veces los mensajes reales (Documentos, Asistencia Humanitaria y Ninguna) o 10 veces (el resto),",
        "contando solo los mensajes que no son de prueba, con la etiqueta final.",
        "",
        "Categoría objetivo = la categoría para la que se generó el mensaje.",
        "Etiqueta LLM = la categoría que el LLM le pone al releerlo, sin contexto.",
        "Revisar = sí: la Etiqueta LLM no coincide con la Categoría objetivo. Esas filas van al final y sombreadas.",
        "Para entrenar: usar la Categoría objetivo de las filas donde Revisar está vacío (objetivo y LLM coinciden).",
        "Categoría corregida: dejar vacío para mantener la categoría objetivo; elegir otra categoría para moverlo.",
        "Nada se mueve ni se elimina automáticamente.",
        "",
        "RESUMEN (objetivo / generados / conservados tras quitar duplicados)",
        *[f"{r.Categoría}: {r.objetivo} / {r.generados} / {r.conservados}" + (f"  (faltan {r.objetivo - r.conservados})" if r.conservados < r.objetivo else "")
          for r in table.itertuples()],
        f"Total conservados: {len(syn)}"]
    book = Workbook()
    book.remove(book.active)
    ws = book.create_sheet("Guia_revision")
    for i, line in enumerate(guia, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 120
    lists = book.create_sheet("_listas")
    for i, v in enumerate(CLASSES, 1):
        lists.cell(i, 1, v)
    lists.sheet_state = "hidden"
    ws = book.create_sheet("Sinteticos")
    widths = {"Mensaje sintético": 70, "Categoría objetivo": 34, "Etiqueta LLM": 34, "Categoría corregida": 34,
              "Comentario": 30, "Revisar": 10, "id": 8}
    for j, h in enumerate(df.columns, 1):
        ws.cell(1, j, h).font = Font(bold=True)
        ws.column_dimensions[ws.cell(1, j).column_letter].width = widths[h]
    yellow = PatternFill("solid", fgColor="FFF2CC")
    flag_col = list(df.columns).index("Revisar")
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, x in enumerate(row, 1):
            if x is not None and not pd.isna(x) and x != "":
                ws.cell(i, j, x.item() if hasattr(x, "item") else x)
        if row[flag_col] == "sí":
            for j in range(1, len(df.columns) + 1):
                ws.cell(i, j).fill = yellow
    col = ws.cell(1, list(df.columns).index("Categoría corregida") + 1).column_letter
    dv = DataValidation(type="list", formula1=f"=_listas!$A$1:$A${len(CLASSES)}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{len(df) + 1}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(df.columns)).column_letter}{len(df) + 1}"
    book.move_sheet("_listas", offset=1)
    book.save(path)


def _load_work(wd: Path) -> pd.DataFrame:
    f = wd / "work.parquet"
    if not f.exists():
        v1 = work_dir("v1") / "work.parquet"
        if wd != v1.parent and v1.exists():          # a new version starts from the same ids and split
            wd.mkdir(parents=True, exist_ok=True)
            shutil.copy(v1, f)
        else:
            raise SystemExit(f"{f} not found. Run `prep` first.")
    return pd.read_parquet(f)


def _corrections(wd: Path) -> pd.DataFrame | None:
    f = wd / "corrections.parquet"
    return pd.read_parquet(f) if f.exists() else None


def _cmd_prep(a, wd):
    from . import redact
    df = cz.load_labeled()
    ner = redact.load_ner()
    work = prepare(df, scrub=lambda t: redact.scrub(t, nlp=ner))
    wd.mkdir(parents=True, exist_ok=True)
    work.to_parquet(wd / "work.parquet")
    t = work[work["split"] == "test"]
    assert len(t) == 588 and (t["label"] != NONE).sum() == 501, "frozen test set changed"
    print(f"{len(work):,} messages · dropped by the scrub (never sent): {int(work['text'].isna().sum())}")
    print(work["split"].value_counts().to_dict(), "| test with text:", int(t["text"].notna().sum()))


def _write_prompts(work, wd, hard=True, control=False):
    system = build_label_prompt(work, corrections=None if control else _corrections(wd), hard=hard)
    assert_no_test_text(system, work)
    (wd / ("prompt_label.txt" if hard else "prompt_label_nohard.txt")).write_text(system, encoding="utf-8")
    return system


def _cmd_pilot(a, wd):
    work = _load_work(wd)
    system = _write_prompts(work, wd)
    lab = work[work["label"].notna() & work["text"].notna()]
    rng = np.random.default_rng(5)
    pick = np.concatenate([rng.choice(lab.index[lab["split"] == s], a.n // 2, replace=False) for s in ("test", "dev")])
    rows = work.loc[pick].set_index("row").sample(frac=1, random_state=5)
    jobs = label_jobs(rows, "p_")
    print(f"pilot: {len(rows)} messages ({a.n // 2} blind test + {a.n // 2} dev) in {len(jobs)} batches")
    if a.dry_run:
        return
    run_batches(jobs, wd / "pilot_batches", LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)
    pred = read_labels(wd / "pilot_batches")
    w = work.set_index("row").loc[list(pred)]
    for s in ("test", "dev"):
        k = w[w["split"] == s]
        if len(k):
            r = agreement(k["label"], [pred[i] for i in k.index])
            print(f"{s}: n={r['n']} accuracy={r['accuracy']:.3f} kappa={r['kappa']:.3f} Ninguna={r['share_ninguna']:.3f}")


def _cmd_label(a, wd):
    work = _load_work(wd)
    system = _write_prompts(work, wd)
    rows = work[work["text"].notna()].set_index("row").sample(frac=1, random_state=31)
    jobs = label_jobs(rows, "l_")
    print(f"label: {len(rows)} messages in {len(jobs)} batches")
    if not a.dry_run:
        run_batches(jobs, wd / "label_batches", LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)


def _embed(texts):
    from . import nlp
    return nlp.embed_documents(texts, batch_size=16)


def _cmd_synth(a, wd):
    work = _load_work(wd)
    jobs = synth_jobs(work, {c: SYNTH_PER_CLASS for c in CLASSES}, "r0_")
    print(f"synth: {len(jobs)} calls for {SYNTH_PER_CLASS} messages in each of {len(CLASSES)} classes")
    if a.dry_run:
        return
    syn = build_synthetic(work, wd / "synth_batches", LLMClient(), SYNTH_SYSTEM, embed=_embed, workers=a.workers)
    collect_synth(wd / "synth_batches").to_csv(wd / "synthetic_raw.csv", index=False)
    syn.to_csv(wd / "synthetic.csv", index=False)
    print(syn["category"].value_counts().to_dict())


def _cmd_check(a, wd):
    work = _load_work(wd)
    system = _write_prompts(work, wd)
    syn = pd.read_csv(wd / "synthetic.csv", keep_default_na=False).sample(frac=1, random_state=37)
    syn["context"] = ""
    jobs = label_jobs(syn, "s_", with_context=False)
    print(f"check: {len(syn)} synthetic messages in {len(jobs)} batches")
    if not a.dry_run:
        run_batches(jobs, wd / "synth_check_batches", LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)


def _cmd_report(a, wd):
    work = _load_work(wd)
    pred = read_labels(wd / "label_batches")
    syn = pd.read_csv(wd / "synthetic.csv", keep_default_na=False)
    syn_pred = read_labels(wd / "synth_check_batches")
    w = work.assign(llm=work["row"].map(pred), client=client_label(work))
    summary = {"Mensajes etiquetados por el LLM": int(w["llm"].notna().sum()),
               "Mensajes anonimizados que no se enviaron": int(work["text"].isna().sum()),
               "Sintéticos": len(syn),
               "Sintéticos que el LLM no pone en su categoría": int(sum(syn_pred.get(i) != c for i, c in enumerate(syn["category"]) if i in syn_pred))}
    for s, name in (("test", "prueba"), ("dev", "entrenamiento")):
        k = w[(w["split"] == s) & w["llm"].notna() & w["client"].notna()]
        r = agreement(k["client"], k["llm"])
        summary[f"Acuerdo LLM-cliente en {name}"] = f"{r['accuracy']:.1%} (n={r['n']}, kappa {r['kappa']:.2f})"
        print(f"{s}: n={r['n']} accuracy={r['accuracy']:.3f} kappa={r['kappa']:.3f}")
        print(r["report"].round(2).to_string())
    path = datasets.DATASETS_DIR / "labels" / "messages_category_revision_haiku.xlsx"
    write_review_workbook(path, work, pred, syn, syn_pred, summary)
    print(f"workbook: {path.name}")


def _cmd_corr(a, wd):
    """Ingest the reviewer's corrections (no model calls) and report who the reviewer agreed with."""
    work = _load_work(wd)
    c = load_corrections(a.corrections, work)
    c.to_parquet(wd / "corrections.parquet")
    print(f"{len(c)} corrections · usable for the prompt (dev): {int((c['split'] == 'dev').sum())} · "
          f"held out (test): {int((c['split'] == 'test').sum())}")
    print(pd.crosstab(c["split"], c["verdict"]).to_string())
    dup = c[c["conflict"]]
    print(f"texts corrected to two different classes: {dup['text'].nunique()} (ids {sorted(dup['row'])}); "
          "never used as few-shots")


def _v1_pred() -> dict[int, str]:
    return read_labels(work_dir("v1") / "label_batches")


def _cmd_baseline(a, wd):
    """Haiku v1 against the gold set, before anything changes (no model calls)."""
    work, c = _load_work(wd), _corrections(wd)
    if c is None:
        raise SystemExit("Run `corr` first.")
    print(score_gold(gold_table(work, _v1_pred(), c), _v1_pred()).round(3).to_string(index=False))


def _cmd_devcheck(a, wd):
    """Gate: relabel the dev corrections plus a sample of agreed dev rows with the corrected prompt.

    The messages are always the same ones (the sample excludes the hard-first shots), so runs compare.
    `--no-hard` shows random corrected examples instead of the hard-first ones and writes its own folder;
    scoring then also leaves out that variant's own shots.
    """
    work, c = _load_work(wd), _corrections(wd)
    if c is None:
        raise SystemExit("Run `corr` first.")
    system = _write_prompts(work, wd, hard=not a.no_hard, control=a.control)
    gold = gold_table(work, _v1_pred(), c)
    gold = gold[(gold["split"] == "dev") & ~gold["row"].isin(shot_rows(work, c))]      # never score a shot on itself
    rng = np.random.default_rng(7)
    agreed = gold[gold["group"] == "agreed"]
    pick = pd.concat([gold[gold["group"] != "agreed"],
                      agreed.loc[rng.choice(agreed.index, min(a.n, len(agreed)), replace=False)]])
    rows = work.set_index("row").loc[pick["row"]].sample(frac=1, random_state=7)
    jobs = label_jobs(rows, "d_")
    folder = wd / ("devcheck_control_batches" if a.control else
                   "devcheck_nohard_batches" if a.no_hard else "devcheck_batches")
    print(f"devcheck{' (no hard-first shots)' if a.no_hard else ''}: {len(rows)} messages in {len(jobs)} batches")
    if a.dry_run:
        return
    run_batches(jobs, folder, LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)
    runs = {"v1": _v1_pred(), "v2": read_labels(folder)}
    other = wd / "devcheck_batches"
    if (a.no_hard or a.control) and other.exists():
        runs["v2 hard-first"] = read_labels(other)
    if a.control:                                   # the v1 prompt again: the run-to-run noise floor
        runs["v1 prompt rerun"] = runs.pop("v2")
    elif a.no_hard:                                 # keep the comparison fair: drop this variant's own shots too
        pick = pick[~pick["row"].isin(shot_rows(work, c, hard=False))]
    ok = lambda p: pick["gold"] == pick["row"].map(p)     
    for name, p in runs.items():
        print(f"{name}:"); print(score_gold(pick, p).round(3).to_string(index=False))
    new = runs.get("v2", runs.get("v1 prompt rerun"))
    print(f"fixed {int((~ok(runs['v1']) & ok(new)).sum())} · broken {int((ok(runs['v1']) & ~ok(new)).sum())}")


def _cmd_testcheck(a, wd):
    """The honest comparison: the 588 frozen-test messages against the client and the 66 corrections.

    `--control` = the v1 examples with the current rules; otherwise the corrected examples. `--tag` names the
    run (each tag is its own folder, so repeats measure the run-to-run noise). Every finished run found in the
    work folder is scored, next to Haiku v1's single run.
    """
    work, c = _load_work(wd), _corrections(wd)
    if c is None:
        raise SystemExit("Run `corr` first.")
    kind = "control" if a.control else "v2"
    system = _write_prompts(work, wd, control=a.control)
    rows = work[(work["split"] == "test") & work["text"].notna()].set_index("row").sample(frac=1, random_state=11)
    jobs = label_jobs(rows, "t_")
    folder = wd / f"testcheck_{kind}_{a.tag}_batches"
    print(f"testcheck {kind}/{a.tag}: {len(rows)} blind test messages in {len(jobs)} batches")
    if a.dry_run:
        return
    run_batches(jobs, folder, LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)
    runs = {"v1 (first prompt)": _v1_pred()} | {f.name.removeprefix("testcheck_").removesuffix("_batches"): read_labels(f)
                                                 for f in sorted(wd.glob("testcheck_*_batches"))}
    k = work.assign(client=client_label(work))
    k = k[k["row"].isin(runs[next(r for r in runs if r != "v1 (first prompt)")]) & k["client"].notna()]
    gold = gold_table(work, runs["v1 (first prompt)"], c)
    gold = gold[(gold["split"] == "test") & (gold["group"] != "agreed")]            # the 66 reviewed test messages
    print(f"{'run':<22}{'vs client':>10}{'kappa':>7}{'Ninguna':>9}{'vs 66 reviewed':>16}{'client-right (4)':>18}")
    for name, p in runs.items():
        r = agreement(k["client"], k["row"].map(p))
        g = score_gold(gold, p).set_index("subset")["accuracy"]
        print(f"{name:<22}{r['accuracy']:>10.3f}{r['kappa']:>7.2f}{r['share_ninguna']:>9.3f}"
              f"{g['all']:>16.3f}{g.get('group=client', float('nan')):>18.3f}")


def _cmd_compare(a, wd):
    """How many labels the corrected prompt changed, by class and by client-labelled vs Haiku-only."""
    work = _load_work(wd)
    old, new = _v1_pred(), read_labels(wd / "label_batches")
    w = work.assign(client=client_label(work), v1=work["row"].map(old), v2=work["row"].map(new))
    w = w[w["v1"].notna() & w["v2"].notna()]
    ch = w[w["v1"] != w["v2"]]
    print(f"{len(w):,} compared · changed {len(ch):,} ({len(ch) / len(w):.1%})")
    print(pd.crosstab(ch["v1"], ch["v2"]).to_string())
    print("changed, client-labelled:", int(ch["client"].notna().sum()), "· Haiku-only:", int(ch["client"].isna().sum()))


def _cmd_table(a, wd):
    work = _load_work(wd)
    t = labels_table(work, read_labels(wd / "label_batches"))
    path = datasets.DATASETS_DIR / "labels" / "messages_category_haiku_all.xlsx"
    write_labels_xlsx(path, t)
    print(f"{path.name}: {len(t):,} rows · agree with client {(t['Coincide cliente'] == 'sí').sum():,} of "
          f"{(t['Coincide cliente'] != '').sum():,}")


def _cmd_final(a, wd):
    """The v1 Haiku labels with the reviewer's corrections on top (no model calls)."""
    v1, v2 = work_dir("v1"), work_dir("v2")
    c = _corrections(v2)
    if c is None:
        raise SystemExit("Run `--version v2 corr` first.")
    t = labels_table(_load_work(v1), read_labels(v1 / "label_batches"), c)
    path = datasets.DATASETS_DIR / "labels" / "messages_category_haiku_final.xlsx"
    write_labels_xlsx(path, t)
    print(f"{path.name}: {len(t):,} rows")
    print(t["Fuente final"].value_counts().to_string())


FINAL_XLSX = "messages_category_haiku_final.xlsx"
SYNTH_XLSX = "messages_category_synthetic_haiku.xlsx"


def _final_work(wd):
    work = _load_work(wd)
    return load_final(datasets.DATASETS_DIR / "labels" / FINAL_XLSX, work)


def _cmd_synth2(a, wd):
    """Synthetic messages sized from the final labels: 3x / 10x the real non-test messages per class."""
    wf = _final_work(wd)
    targets, pools = synth_targets(wf), synth_pools(wf)
    if a.only:
        targets = {c: n for c, n in targets.items() if c == a.only}
    thin = {c: len(pools[c]) for c in targets if len(pools[c]) < 30}
    if thin:
        raise SystemExit(f"too few seed messages: {thin}")
    print("class target (seed pool):")
    for c, n in targets.items():
        print(f"  {c}: {n:,} ({len(pools[c]):,})")
    jobs = synth_jobs(wf, targets, "r0_", pools=pools)
    print(f"synth2: {sum(targets.values()):,} messages in {len(jobs)} round-0 calls (no test text in any prompt)")
    if a.dry_run:
        return
    embed = cached_embedder(wd / "embeddings.pkl", _embed)
    syn = build_synthetic(wf, wd / "synth_batches", LLMClient(), SYNTH_SYSTEM, embed=embed, targets=targets,
                          pools=pools, workers=a.workers)
    raw = collect_synth(wd / "synth_batches")
    raw.to_csv(wd / "synthetic_raw.csv", index=False)
    syn.to_csv(wd / "synthetic.csv", index=False)
    got, made = syn["category"].value_counts(), raw["category"].value_counts()
    print("class: target / generated / kept")
    for c, n in targets.items():
        print(f"  {c}: {n:,} / {made.get(c, 0):,} / {got.get(c, 0):,}" + ("  SHORT" if got.get(c, 0) < n else ""))


def _cmd_check2(a, wd):
    """Re-label every synthetic message with the v1 labelling prompt, no context."""
    work = _load_work(wd)
    system = _write_prompts(work, wd, control=True)           # control = no corrections = the v1 prompt
    syn = pd.read_csv(wd / "synthetic.csv", keep_default_na=False).sample(frac=1, random_state=37)
    syn["context"] = ""
    jobs = label_jobs(syn, "s_", with_context=False)
    print(f"check2: {len(syn):,} synthetic messages in {len(jobs)} batches")
    if not a.dry_run:
        run_batches(jobs, wd / "synth_check_batches", LLMClient(), system, LABEL_SCHEMA, check_labels, workers=a.workers)


def _cmd_synthreport(a, wd):
    wf = _final_work(wd)
    targets = synth_targets(wf)
    syn = pd.read_csv(wd / "synthetic.csv", keep_default_na=False)
    raw = pd.read_csv(wd / "synthetic_raw.csv", keep_default_na=False)
    pred = read_labels(wd / "synth_check_batches")
    table = pd.DataFrame({"Categoría": CLASSES, "objetivo": [targets[c] for c in CLASSES],
                          "generados": [int((raw["category"] == c).sum()) for c in CLASSES],
                          "conservados": [int((syn["category"] == c).sum()) for c in CLASSES]})
    path = datasets.DATASETS_DIR / "labels" / SYNTH_XLSX
    write_synthetic_workbook(path, syn, pred, table)
    flagged = sum(pred.get(i) not in (None, c) for i, c in enumerate(syn["category"]))
    print(table.to_string(index=False))
    print(f"{path.name}: {len(syn):,} rows · flagged Revisar = sí: {flagged:,} · not yet checked: {len(syn) - len(pred):,}")


def main(argv=None) -> None:
    p = argparse.ArgumentParser(prog="python -m sami.llm_labels", description=__doc__.split("\n\n")[0])
    p.add_argument("stage", choices=["prep", "pilot", "label", "synth", "check", "report", "table",
                                     "corr", "baseline", "devcheck", "testcheck", "compare", "final",
                                     "synth2", "check2", "synthreport"])
    p.add_argument("--version", default="v1", help="work folder version: v1 = llm_haiku, v2 = llm_haiku_v2")
    p.add_argument("--no-hard", action="store_true", help="devcheck: random corrected few-shots, not hard-first")
    p.add_argument("--only", help="synth2: generate just this class (the small-class checkpoint)")
    p.add_argument("--tag", default="r1", help="testcheck: name of the run, so repeats do not overwrite")
    p.add_argument("--control", action="store_true", help="devcheck/testcheck: the v1 examples (no corrections) with the current rules")
    p.add_argument("--corrections", type=Path,
                   default=datasets.DATASETS_DIR / "labels" / "messages_category_revised_haiku.xlsx",
                   help="reviewed workbook read by `corr`")
    p.add_argument("--dry-run", action="store_true", help="build prompts and job lists, call no model")
    p.add_argument("--workers", type=int, default=WORKERS, help="parallel model calls (default %(default)s)")
    p.add_argument("-n", type=int, default=300, help="pilot size (half blind test, half dev); agreed rows in devcheck")
    a = p.parse_args(argv)
    wd = work_dir(a.version)
    {"prep": _cmd_prep, "pilot": _cmd_pilot, "label": _cmd_label, "synth": _cmd_synth,
     "check": _cmd_check, "report": _cmd_report, "table": _cmd_table, "corr": _cmd_corr,
     "baseline": _cmd_baseline, "devcheck": _cmd_devcheck, "testcheck": _cmd_testcheck,
     "compare": _cmd_compare, "final": _cmd_final, "synth2": _cmd_synth2,
     "check2": _cmd_check2, "synthreport": _cmd_synthreport}[a.stage](a, wd)


if __name__ == "__main__":
    main()
